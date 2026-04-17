import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import io
from decimal import Decimal
import openpyxl

from backend.ocr.models import ExportRow
from backend.export.excel import build_xlsx, COLUMNS
from backend.export.csv_export import build_csv


SAMPLE_ROWS = [
    ExportRow(
        date="2024-03-15",
        category="Meals & Entertainment",
        description="Coffee, Muffin",
        vendor="Starbucks",
        gst=Decimal("0.35"),
        total=Decimal("7.35"),
        base_amount=Decimal("7.00"),
    ),
    ExportRow(
        date="2024-03-16",
        category="Unclassified",
        description="See receipt",
        vendor="Unknown Store",
        gst=None,
        total=Decimal("25.00"),
        base_amount=None,
    ),
]


class TestBuildXlsx:
    def test_returns_bytes(self):
        result = build_xlsx(SAMPLE_ROWS)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_valid_xlsx(self):
        result = build_xlsx(SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        assert "Receipts" in wb.sheetnames

    def test_header_row(self):
        result = build_xlsx(SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        headers = [ws.cell(1, i + 1).value for i in range(len(COLUMNS))]
        assert headers == COLUMNS

    def test_data_rows(self):
        result = build_xlsx(SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.cell(2, 2).value == "Meals & Entertainment"
        assert ws.cell(2, 4).value == "Starbucks"
        assert ws.cell(3, 2).value == "Unclassified"

    def test_money_values(self):
        result = build_xlsx(SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.cell(2, 5).value == pytest.approx(0.35)
        assert ws.cell(2, 6).value == pytest.approx(7.35)
        assert ws.cell(2, 7).value == pytest.approx(7.00)

    def test_empty_rows(self):
        result = build_xlsx([])
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.max_row == 1  # only header


class TestBuildCsv:
    def test_returns_string(self):
        result = build_csv(SAMPLE_ROWS)
        assert isinstance(result, str)

    def test_header_line(self):
        result = build_csv(SAMPLE_ROWS)
        first_line = result.splitlines()[0]
        assert "date" in first_line
        assert "category" in first_line
        assert "gst" in first_line

    def test_data_lines(self):
        result = build_csv(SAMPLE_ROWS)
        lines = result.splitlines()
        assert len(lines) == 3  # header + 2 rows
        assert "Starbucks" in lines[1]
        assert "Meals & Entertainment" in lines[1]

    def test_missing_money_is_empty(self):
        result = build_csv(SAMPLE_ROWS)
        lines = result.splitlines()
        fields = lines[2].split(",")
        # base_amount is None for second row → empty string
        assert fields[-1] == ""


import pytest
