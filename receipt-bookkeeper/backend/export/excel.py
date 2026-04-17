from io import BytesIO
from decimal import Decimal
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..ocr.models import ExportRow

COLUMNS = [
    "Date",
    "Category",
    "Description",
    "Vendor",
    "GST",
    "Total (Incl. GST)",
    "Base Amount",
]

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2563EB")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
MONEY_FORMAT = '$#,##0.00'
UNCLASSIFIED_FILL = PatternFill("solid", fgColor="FEF3C7")


def _dec(value: Optional[Decimal]) -> Optional[float]:
    return float(value) if value is not None else None


def build_xlsx(rows: list[ExportRow]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receipts"

    # Header row
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 22

    # Data rows
    for row_idx, row in enumerate(rows, 2):
        values = [
            str(row.date) if row.date else "",
            row.category,
            row.description,
            row.vendor or "",
            _dec(row.gst),
            _dec(row.total),
            _dec(row.base_amount),
        ]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center")

        # Highlight unclassified rows in amber
        if row.category == "Unclassified":
            for col_idx in range(1, len(COLUMNS) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = UNCLASSIFIED_FILL

        # Money format for GST, Total, Base Amount (columns 5, 6, 7)
        for col_idx in (5, 6, 7):
            ws.cell(row=row_idx, column=col_idx).number_format = MONEY_FORMAT

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(
            max_len + 4, 14
        )

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
